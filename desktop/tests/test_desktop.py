from pathlib import Path

import pytest

from pii_roundtrip_desktop import documents, i18n
from pii_roundtrip_desktop.store import Store


def test_docx_keeps_paragraphs_and_tables(tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Договор с Ивановым Иваном Ивановичем")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "СНИЛС"
    table.rows[0].cells[1].text = "112-233-445 95"
    doc.add_paragraph("Подпись")
    p = tmp_path / "d.docx"
    doc.save(p)
    text = documents.read_text(p)
    assert text.splitlines() == ["Договор с Ивановым Иваном Ивановичем", "СНИЛС | 112-233-445 95", "Подпись"]


def test_text_in_cp1251_and_unsupported_formats(tmp_path: Path):
    p = tmp_path / "a.txt"
    p.write_bytes("Иванов".encode("cp1251"))
    assert documents.read_text(p) == "Иванов"
    with pytest.raises(documents.DocumentError):
        documents.read_text(tmp_path / "a.xls")


def test_pdf_without_text_layer_is_reported(tmp_path: Path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    p = tmp_path / "scan.pdf"
    with open(p, "wb") as f:
        w.write(f)
    with pytest.raises(documents.ScannedPDF):
        documents.read_text(p)


def test_word_round_trip(tmp_path: Path):
    p = tmp_path / "out.docx"
    documents.write_docx(p, "first\nsecond")
    assert documents.read_text(p) == "first\nsecond\n"


def test_store_saves_sessions_and_checks_the_password(tmp_path: Path):
    from pii_roundtrip.core import Engine
    from pii_roundtrip.detect import Models

    eng = Engine(Models(use_pii_guard=False, use_gliner=False))
    _, s = eng.anonymize("СНИЛС 112-233-445 95", "ru")
    store = Store("secret", tmp_path)
    e = store.save("договор.docx", s)
    _, s2 = eng.anonymize("Повторно СНИЛС 112-233-445 95", "ru", store.load(e.key))
    e2 = store.save("договор.docx", s2, e.key)  # a follow-up keeps one entry
    assert [x.key for x in store.entries()] == [e.key] == [e2.key]
    assert store.load(e.key).mapping == s.mapping
    assert not Store("wrong", tmp_path).check()
    store.save_texts(e.key, {"request": "[SNILS_1]", "answer": "СНИЛС 112-233-445 95"})
    assert Store("secret", tmp_path).load_texts(e.key)["answer"] == "СНИЛС 112-233-445 95"
    raw = (tmp_path / (e.key + ".texts")).read_bytes()
    assert "112-233".encode() not in raw  # the texts file is encrypted
    store.delete(e.key)
    assert store.entries() == []
    assert (tmp_path.stat().st_mode & 0o777) == 0o700


def test_language_and_summary():
    assert i18n.detect_language("Договор аренды, Mr Smith") == "ru"
    assert i18n.detect_language("Tenancy agreement with Иван") == "en-GB"
    t = i18n.TEXTS["ru"]
    assert i18n.summary({"[PERSON_1]": "a", "[PERSON_2]": "b", "[SNILS_1]": "c"}, t) == "имена 2 · номера 1"
    assert i18n.mixed("Договор с John Smith Limited London") and not i18n.mixed("Договор аренды")


def test_clean_and_restore_work_with_the_network_forbidden(tmp_path: Path, monkeypatch):
    import socket

    from pii_roundtrip_desktop.app import forbid_network
    from pii_roundtrip_desktop.engine import AppEngine

    for name in ("connect", "connect_ex"):
        monkeypatch.setattr(socket.socket, name, getattr(socket.socket, name))
    monkeypatch.setattr(socket, "create_connection", socket.create_connection)
    forbid_network()
    with pytest.raises(OSError):
        socket.create_connection(("example.com", 443))
    eng = AppEngine()
    eng.load()  # models from the local cache, or rules only: never the network
    clean, session = eng.anonymize("СНИЛС 112-233-445 95, карта 4111 1111 1111 1111", "ru", None)
    assert "112-233" not in clean
    assert eng.deanonymize(clean, session) == "СНИЛС 112-233-445 95, карта 4111 1111 1111 1111"


def test_spans_match_the_engine():
    from pii_roundtrip.core import Engine
    from pii_roundtrip.detect import Models
    from pii_roundtrip_desktop.engine import AppEngine

    app = AppEngine()
    app._engine = Engine(Models(use_pii_guard=False, use_gliner=False))
    text = "СНИЛС 112-233-445 95, карта 4111 1111 1111 1111, повторно 112-233-445 95."
    clean, session, spans = app.anonymize_with_spans(text, "ru", None)
    ref_clean, ref_session = app._engine.anonymize(text, "ru")
    assert clean == ref_clean and session.mapping == ref_session.mapping
    rebuilt, pos = [], 0
    for start, end, label in spans:
        rebuilt += [text[pos:start], label]
        pos = end
    assert "".join(rebuilt) + text[pos:] == clean


def _letter(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].header.paragraphs[0].text = "ООО «Восход», директор Сидорова А.В."
    p = doc.add_paragraph("Договор с ")
    p.add_run("Ивановым Петром").bold = True
    p.add_run(", СНИЛС 112-233-445 95.")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Карта"
    t.rows[0].cells[1].text = "4111 1111 1111 1111"
    doc.core_properties.author = "Сидорова Анна"
    doc.save(path)


def test_clean_word_copy_keeps_formatting_and_cleans_headers(tmp_path: Path):
    from docx import Document

    src = tmp_path / "letter.docx"
    _letter(src)
    text = documents.read_text(src)
    assert text.startswith("ООО «Восход», директор Сидорова А.В.")
    spans = []
    for value, label in (
        ("Сидорова А.В.", "[PERSON_1]"),
        ("Ивановым Петром", "[PERSON_2]"),
        ("112-233-445 95", "[SNILS_1]"),
        ("4111 1111 1111 1111", "[CARD_1]"),
    ):
        i = text.index(value)
        spans.append((i, i + len(value), label))
    dst = tmp_path / "letter.clean.docx"
    documents.save_clean_docx(
        src,
        dst,
        sorted(spans),
        ["Сидорова А.В.", "Ивановым Петром", "112-233-445 95", "4111 1111 1111 1111", "Сидорова Анна"],
    )
    out = Document(dst)
    body = out.paragraphs[0]
    assert body.text == "Договор с [PERSON_2], СНИЛС [SNILS_1]."
    assert any(r.bold and r.text == "[PERSON_2]" for r in body.runs)  # formatting kept
    assert out.sections[0].header.paragraphs[0].text == "ООО «Восход», директор [PERSON_1]"
    assert out.tables[0].rows[0].cells[1].text == "[CARD_1]"
    assert out.core_properties.author == ""


def test_clean_word_copy_is_refused_if_a_value_remains(tmp_path: Path):
    src = tmp_path / "letter.docx"
    _letter(src)
    with pytest.raises(documents.UnsafeCopy):
        documents.save_clean_docx(src, tmp_path / "out.docx", [], ["Ивановым Петром"])
    assert not (tmp_path / "out.docx").exists()


def test_pdf_and_text_output(tmp_path: Path):
    documents.write_pdf(tmp_path / "a.pdf", "Договор с [PERSON_1]\n\nСНИЛС [SNILS_1]")
    assert documents.read_text(tmp_path / "a.pdf").startswith("Договор с [PERSON_1]")
    documents.write_text(tmp_path / "a.txt", "текст")
    assert documents.read_text(tmp_path / "a.txt") == "текст"


def test_clean_pdf_copy_keeps_pages_and_removes_values(tmp_path: Path):
    src = tmp_path / "in.pdf"
    documents.write_pdf(src, "Договор с Ивановым Петром\n\nСНИЛС 112-233-445 95 и прочий текст договора для длины")
    text = documents.read_text(src)
    spans = []
    for value, label in (("Ивановым Петром", "[PERSON_1]"), ("112-233-445 95", "[SNILS_1]")):
        i = text.index(value)
        spans.append((i, i + len(value), label))
    dst = tmp_path / "out.pdf"
    documents.save_clean_pdf(src, dst, spans, ["Ивановым Петром", "112-233-445 95"])
    out = documents.read_text(dst)
    assert "Ивановым" not in out and "112-233" not in out
    assert "[PERSON_1]" in out and "[SNILS_1]" in out  # the invisible text layer


def test_plain_docx_has_no_author(tmp_path):
    from docx import Document

    path = tmp_path / "plain.docx"
    documents.write_docx(path, "[PERSON_1] подписал договор.")
    props = Document(str(path)).core_properties
    assert props.author == "" and props.comments == "" and props.last_modified_by == ""


def test_restore_keeps_case_across_lines():
    from pii_roundtrip_desktop.engine import AppEngine

    e = AppEngine()
    text = "г. Москва\nИванов Пётр Сергеевич, паспорт 45 10 123456, именуемый Арендодатель."
    clean, s = e.anonymize(text, "ru", None)
    back = e.deanonymize(clean + "\nДоговор подписан с [PERSON_1].", s)
    assert back.startswith(text)
    assert back.endswith("с Ивановым Петром Сергеевичем.")


def test_replacement_spans_follow_placeholders():
    before = "Договор с [PERSON_1], паспорт [PASSPORT_RU_1]. [ORG_1]"
    after = "Договор с Ивановым Петром, паспорт 45 10 123456. ORACLE Ltd"
    spans = documents.replacement_spans(before, after)
    assert [before[a:b] for a, b, _ in spans] == ["[PERSON_1]", "[PASSPORT_RU_1]", "[ORG_1]"]
    assert [v for _, _, v in spans] == ["Ивановым Петром", "45 10 123456", "ORACLE Ltd"]


def test_pdf_clean_then_restore_in_place(tmp_path):
    src = tmp_path / "letter.pdf"
    documents.write_pdf(src, "Арендатор Иванов Пётр Сергеевич\nСНИЛС 112-233-445 95\nг. Москва")
    text = documents.read_text(src)
    spans = []
    for value, label in (("Иванов Пётр Сергеевич", "[PERSON_1]"), ("112-233-445 95", "[SNILS_1]")):
        i = text.index(value)
        spans.append((i, i + len(value), label))
    clean = tmp_path / "letter.clean.pdf"
    documents.save_clean_pdf(src, clean, spans, ["Иванов Пётр Сергеевич", "112-233-445 95"])
    labelled = documents.read_text(clean)
    assert "[PERSON_1]" in labelled and "[SNILS_1]" in labelled and "Иванов" not in labelled
    restored = labelled.replace("[PERSON_1]", "Иванов Пётр Сергеевич").replace("[SNILS_1]", "112-233-445 95")
    back = tmp_path / "letter.restored.pdf"
    documents.save_restored_pdf(clean, back, documents.replacement_spans(labelled, restored))
    again = documents.read_text(back)
    assert "Иванов Пётр Сергеевич" in again and "112-233-445 95" in again and "[PERSON_1]" not in again


def test_delete_all_removes_every_vault(tmp_path):
    from pii_roundtrip.core import Session

    store = Store("secret", tmp_path)
    for n in ("a.docx", "b.pdf"):
        e = store.save(n, Session(lang="ru"))
        store.save_texts(e.key, {"request": "x"})
    (tmp_path / "old-unlisted.vault").write_bytes(b"x")
    store.delete_all()
    assert store.entries() == [] and not list(tmp_path.iterdir())


def test_nothing_is_saved_or_copied_while_cleaning(tmp_path, monkeypatch):
    import os

    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    from pii_roundtrip_desktop import app as A

    QApplication.instance() or QApplication([])

    class Ready:  # no model loading in this test
        ready = True

    w = A.MainWindow(Store("secret", tmp_path / "vault"), i18n.TEXTS["ru"], engine=Ready())
    target = tmp_path / "out.txt"
    monkeypatch.setattr(A.QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(target), "")))
    w._working("Иванов Пётр Сергеевич, паспорт 45 10 123456", "draft.txt")  # the card shows the original
    w.save_docx()
    w.copy_all()
    assert not target.exists()
    assert "Иванов" not in (QApplication.clipboard().text() or "")


def test_replacement_spans_survive_a_value_containing_the_next_text():
    # the address itself contains ", г. ", the text that follows its placeholder
    before = "Адрес: [ADDRESS_1], г. Москва\nИНН [INN_1]"
    after = "Адрес: 125047, г. Москва, ул. Гашека, г. Москва\nИНН 7706440687"
    spans = documents.replacement_spans(before, after)
    out, pos = [], 0
    for a, b, v in spans:
        out += [before[pos:a], v]
        pos = b
    out.append(before[pos:])
    assert "".join(out) == after


def test_redrawn_cleaned_pdf_is_not_guessed(tmp_path):
    """Pages redrawn after cleaning: no guessing on a picture; a page of real text is restored."""
    import pypdfium2 as pdfium
    from reportlab.pdfgen import canvas

    src = tmp_path / "letter.pdf"
    documents.write_pdf(src, "Арендатор Иванов Пётр Сергеевич\nг. Москва")
    text = documents.read_text(src)
    i = text.index("Иванов Пётр Сергеевич")
    clean = tmp_path / "clean.pdf"
    documents.save_clean_pdf(src, clean, [(i, i + 21, "[PERSON_1]")], ["Иванов Пётр Сергеевич"])
    keywords = pdfium.PdfDocument(str(clean)).get_metadata_dict()["Keywords"]

    def redraw(path, visible):
        c = canvas.Canvas(str(path))
        c.setKeywords(keywords)  # the record survives, the picture of the page does not
        t = c.beginText(72, 760)
        if not visible:
            t.setTextRenderMode(3)  # a page that is a picture with a hidden text layer
        t.textOut("Arendator [PERSON_1] g. Moskva")
        c.drawText(t)
        c.save()
        labelled = documents.read_text(path)
        return labelled, documents.replacement_spans(labelled, labelled.replace("[PERSON_1]", "Ivanov Petr"))

    picture = tmp_path / "picture.pdf"
    _, spans = redraw(picture, visible=False)
    with pytest.raises(documents.UnplacedLabels):
        documents.save_restored_pdf(picture, tmp_path / "back1.pdf", spans)
    # a page an AI tool wrote anew with real text: the values go where its letters are
    letters = tmp_path / "letters.pdf"
    _, spans = redraw(letters, visible=True)
    documents.save_restored_pdf(letters, tmp_path / "back2.pdf", spans)
    assert "Ivanov Petr" in documents.read_text(tmp_path / "back2.pdf")


def test_reflow_puts_pdf_lines_back_into_paragraphs():
    text = "ДОГОВОР № 1\nг. Москва\nОбщество, имену \nемое Заказчик, заклю\nчили договор о нижесле\nдующем:\n1.1. Исполнитель выполняет рабо\nты в срок.\n1.2. Оплата."
    assert documents.reflow(text) == (
        "ДОГОВОР № 1\nг. Москва\nОбщество, именуемое Заказчик, заключили договор о нижеследующем:\n"
        "1.1. Исполнитель выполняет работы в срок.\n1.2. Оплата."
    )


def test_an_edited_cleaned_file_is_not_taken_for_the_original():
    """An AI tool that edits the text keeps the labels: only the whole text proves the file is
    unchanged, and only then may the original file be handed back."""
    from pii_roundtrip_desktop.app import _same_text

    clean = "Договор с [PERSON_1]: срок до 30 апреля, цена 150 000."
    assert _same_text(clean, "Договор с [PERSON_1]:\nсрок до 30 ап­реля, цена 150 000.")
    assert not _same_text(clean, "Договор с [PERSON_1]: срок до 31 мая, цена 180 000.")
